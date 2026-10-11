"""Automated tests for the /pdf and /word document conversion feature.

Run with:

    venv/bin/python test_conversion.py
"""

import asyncio
import json
import logging
import os
import shutil
import tempfile
import threading
import time
import unittest
import zipfile
from unittest import mock

import docx as python_docx
import httpx
import jwt as pyjwt
import pymupdf

import bot
import converter
from converter import (
    ConversionError,
    convert_pdf_to_word,
    convert_word_to_pdf,
    inspect_pdf,
    inspect_word,
    validate_document,
)

FAKE_PDF_BYTES = b"%PDF-1.7\n1 0 obj\n<< >>\nendobj\ntrailer\n<< >>\n%%EOF\n"


# ---------------------------------------------------------------- fixtures


def make_pdf(path: str, text: str = "Halo dunia, ini dokumen uji.", password: str | None = None) -> str:
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 100), text, fontsize=14)
    if password:
        doc.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw=password, owner_pw=password)
    else:
        doc.save(path)
    doc.close()
    return path


def make_image_only_pdf(path: str) -> str:
    """PDF tanpa teks (menyerupai hasil scan)."""
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.draw_rect(pymupdf.Rect(50, 50, 400, 300), color=(0, 0, 0), fill=(0.9, 0.9, 0.9))
    doc.save(path)
    doc.close()
    return path


def make_empty_pdf(path: str) -> str:
    """PDF valid tanpa halaman (pymupdf menolak menyimpan dokumen 0 halaman)."""
    raw = (
        b"%PDF-1.4\n"
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n"
        b"2 0 obj << /Type /Pages /Kids [] /Count 0 >> endobj\n"
        b"trailer << /Root 1 0 R /Size 3 >>\n"
        b"%%EOF\n"
    )
    with open(path, "wb") as fh:
        fh.write(raw)
    return path


def make_docx(path: str, text: str = "Isi dokumen Word uji.") -> str:
    document = python_docx.Document()
    document.add_paragraph(text)
    document.save(path)
    return path


def leftover_tmp_dirs() -> list[str]:
    if not os.path.isdir(bot.DOWNLOAD_DIR):
        return []
    return sorted(name for name in os.listdir(bot.DOWNLOAD_DIR) if name.startswith("convert_"))


def patched_client(transport):
    """Ganti httpx.AsyncClient dengan klien yang memakai transport mock (tanpa rekursi)."""
    real_client = httpx.AsyncClient

    def factory(**kwargs):
        kwargs.pop("transport", None)
        return real_client(transport=transport, **kwargs)

    return mock.patch("httpx.AsyncClient", factory)


# ---------------------------------------------------------------- telegram fakes


class FakeUser:
    def __init__(self, user_id: int):
        self.id = user_id


class FakeDocument:
    def __init__(self, file_name, mime_type, file_size, file_id="file-1"):
        self.file_name = file_name
        self.mime_type = mime_type
        self.file_size = file_size
        self.file_id = file_id


class FakeFile:
    def __init__(self, source_path):
        self.source_path = source_path

    async def download_to_drive(self, custom_path=None, **kwargs):
        if self.source_path is None:
            open(custom_path, "wb").close()  # file kosong
        else:
            shutil.copyfile(self.source_path, custom_path)


class FakeBot:
    def __init__(self, files):
        self.files = files  # file_id -> path lokal

    async def get_file(self, file_id):
        if file_id not in self.files:
            raise RuntimeError("file_id tidak dikenal")
        return FakeFile(self.files[file_id])


class FakeContext:
    def __init__(self, files):
        self.bot = FakeBot(files)


class FakeMessage:
    """Meniru bagian dari telegram.Message yang dipakai handler konversi."""

    def __init__(self, text=None, caption=None, document=None, reply_to_message=None):
        self.text = text
        self.caption = caption
        self.document = document
        self.reply_to_message = reply_to_message
        self.replies = []
        self.edits = []
        self.deleted = False
        self.fail_status = False  # tiru pesan status yang gagal diedit/dihapus
        self.filename = None
        self.caption_sent = None
        self.data = None

    async def reply_text(self, text, **kwargs):
        child = FakeMessage(text=text)
        child.fail_status = self.fail_status
        self.replies.append(child)
        return child

    async def reply_document(self, document=None, filename=None, caption=None, **kwargs):
        child = FakeMessage()
        child.filename = filename
        child.caption_sent = caption
        child.data = document.read()
        self.replies.append(child)
        return child

    async def edit_text(self, text, **kwargs):
        if self.fail_status:
            raise RuntimeError("edit failed")
        self.edits.append(text)
        self.text = text
        return self

    async def delete(self):
        if self.fail_status:
            raise RuntimeError("delete failed")
        self.deleted = True

    @property
    def sent_documents(self):
        return [reply for reply in self.replies if reply.data is not None]


class FakeUpdate:
    def __init__(self, message, user_id: int = 111):
        self.message = message
        self.effective_user = FakeUser(user_id)


PDF_MIME = "application/pdf"
TEST_PUBLIC_KEY = "public-test-key-0123456789abcdef0123"
TEST_SECRET_KEY = "secret-test-key-0123456789abcdef0123"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


# ---------------------------------------------------------------- tests


class ValidationTests(unittest.TestCase):
    """Skenario validasi: ekstensi, MIME type, ukuran."""

    def test_accepts_pdf(self):
        self.assertEqual(validate_document("dokumen.pdf", PDF_MIME, 1000, "pdf"), ".pdf")

    def test_accepts_docx(self):
        self.assertEqual(
            validate_document("dokumen.docx", DOCX_MIME, 1000, "word"), ".docx"
        )

    def test_accepts_legacy_doc(self):
        self.assertEqual(validate_document("dokumen.doc", "application/msword", 1000, "word"), ".doc")

    def test_accepts_generic_mime(self):
        # Telegram kadang mengirim application/octet-stream
        self.assertEqual(validate_document("dokumen.pdf", "application/octet-stream", 10, "pdf"), ".pdf")

    def test_rejects_image_for_pdf(self):
        with self.assertRaises(ConversionError) as ctx:
            validate_document("foto.png", "image/png", 100, "pdf")
        self.assertIn("/pdf", str(ctx.exception))

    def test_rejects_pdf_for_word(self):
        with self.assertRaises(ConversionError) as ctx:
            validate_document("dokumen.pdf", PDF_MIME, 100, "word")
        self.assertIn("/word", str(ctx.exception))

    def test_rejects_missing_filename(self):
        with self.assertRaises(ConversionError):
            validate_document("", PDF_MIME, 100, "pdf")

    def test_rejects_empty_file(self):
        with self.assertRaises(ConversionError) as ctx:
            validate_document("dokumen.pdf", PDF_MIME, 0, "pdf")
        self.assertIn("kosong", str(ctx.exception))

    def test_rejects_oversized_file(self):
        limit = converter.CONVERT_MAX_FILE_SIZE + 1
        with self.assertRaises(ConversionError) as ctx:
            validate_document("dokumen.pdf", PDF_MIME, limit, "pdf")
        self.assertIn("melebihi batas", str(ctx.exception))


class InspectionTests(unittest.TestCase):
    """Deteksi file rusak / terproteksi / kosong sebelum konversi."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="inspect_")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def path(self, name):
        return os.path.join(self.tmp, name)

    def test_valid_pdf_passes(self):
        self.assertTrue(inspect_pdf(make_pdf(self.path("ok.pdf"))))

    def test_corrupted_pdf_rejected(self):
        path = self.path("rusak.pdf")
        with open(path, "wb") as fh:
            fh.write(b"bukan pdf sama sekali " + os.urandom(64))
        with self.assertRaises(ConversionError) as ctx:
            inspect_pdf(path)
        self.assertIn("rusak", str(ctx.exception))

    def test_password_protected_pdf_rejected(self):
        path = self.path("sandi.pdf")
        make_pdf(self.path("_tmp.pdf"), password="rahasia123")
        os.replace(self.path("_tmp.pdf"), path)
        with self.assertRaises(ConversionError) as ctx:
            inspect_pdf(path)
        self.assertIn("password", str(ctx.exception))

    def test_empty_pdf_rejected(self):
        with self.assertRaises(ConversionError) as ctx:
            inspect_pdf(make_empty_pdf(self.path("kosong.pdf")))
        self.assertIn("kosong", str(ctx.exception))

    def test_image_only_pdf_reports_no_text(self):
        self.assertFalse(inspect_pdf(make_image_only_pdf(self.path("scan.pdf"))))

    def test_corrupted_docx_rejected(self):
        path = self.path("rusak.docx")
        with open(path, "wb") as fh:
            fh.write(os.urandom(64))
        with self.assertRaises(ConversionError) as ctx:
            inspect_word(path, ".docx")
        self.assertIn("rusak", str(ctx.exception))

    def test_valid_docx_passes(self):
        inspect_word(make_docx(self.path("ok.docx")), ".docx")  # tidak melempar error

    def test_legacy_doc_is_not_inspected_locally(self):
        path = self.path("lama.doc")
        with open(path, "wb") as fh:
            fh.write(b"\xd0\xcf\x11\xe0" + os.urandom(32))
        inspect_word(path, ".doc")  # dibiarkan dinilai API


class PdfToWordTests(unittest.IsolatedAsyncioTestCase):
    """Konversi PDF -> DOCX secara lokal."""

    async def test_converts_pdf_to_valid_docx(self):
        tmp = tempfile.mkdtemp(prefix="conv_test_")
        self.addCleanup(shutil.rmtree, tmp, True)
        src = make_pdf(os.path.join(tmp, "dokumen.pdf"), text="Paragraf hasil uji konversi")
        dst = os.path.join(tmp, "dokumen.docx")

        await convert_pdf_to_word(src, dst)

        self.assertTrue(zipfile.is_zipfile(dst))
        content = "\n".join(p.text for p in python_docx.Document(dst).paragraphs)
        self.assertIn("Paragraf hasil uji konversi", content)

    async def test_concurrency_is_limited(self):
        tmp = tempfile.mkdtemp(prefix="conv_test_")
        self.addCleanup(shutil.rmtree, tmp, True)
        sources = [make_pdf(os.path.join(tmp, f"doc{i}.pdf")) for i in range(5)]
        original = converter._pdf_to_docx_sync
        state = {"active": 0, "max": 0}
        lock = threading.Lock()

        def tracked(src_path, dst_path):
            with lock:
                state["active"] += 1
                state["max"] = max(state["max"], state["active"])
            try:
                time.sleep(0.05)
                return original(src_path, dst_path)
            finally:
                with lock:
                    state["active"] -= 1

        with mock.patch.object(converter, "_pdf_to_docx_sync", tracked):
            await asyncio.gather(
                *[
                    convert_pdf_to_word(src, os.path.join(tmp, f"out{i}.docx"))
                    for i, src in enumerate(sources)
                ]
            )

        self.assertEqual(state["max"], min(5, converter.CONVERT_MAX_CONCURRENT))
        for i in range(5):
            self.assertTrue(os.path.exists(os.path.join(tmp, f"out{i}.docx")))


class HandlerTests(unittest.IsolatedAsyncioTestCase):
    """Skenario handler /pdf dan /word di bot.py."""

    async def asyncSetUp(self):
        self.tmp = tempfile.mkdtemp(prefix="handler_")
        self.before = leftover_tmp_dirs()

    async def asyncTearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def fixture(self, name, maker, *args, **kwargs):
        return maker(os.path.join(self.tmp, name), *args, **kwargs)

    async def run_reply_command(self, command, document, source_path, user_id=111, files=None):
        """Jalankan /pdf atau /word sebagai balasan ke sebuah dokumen."""
        doc_message = FakeMessage(document=document)
        command_message = FakeMessage(text=f"/{command}", reply_to_message=doc_message)
        update = FakeUpdate(command_message, user_id)
        if files is None:
            files = {document.file_id: source_path}
        context = FakeContext(files)
        handler = bot.pdf_command if command == "pdf" else bot.word_command
        await handler(update, context)
        return command_message

    def assert_no_leftovers(self):
        self.assertEqual(leftover_tmp_dirs(), self.before)

    # --- penggunaan tanpa dokumen

    async def test_pdf_without_document_replies_usage(self):
        message = FakeMessage(text="/pdf")
        await bot.pdf_command(FakeUpdate(message), FakeContext({}))
        self.assertIn("Cara pakai /pdf", message.replies[0].text)
        self.assert_no_leftovers()

    async def test_word_without_document_replies_usage(self):
        message = FakeMessage(text="/word")
        await bot.word_command(FakeUpdate(message), FakeContext({}))
        self.assertIn("Cara pakai /word", message.replies[0].text)
        self.assertIn("iLovePDF", message.replies[0].text)
        self.assert_no_leftovers()

    async def test_reply_to_non_document_replies_usage(self):
        message = FakeMessage(text="/pdf", reply_to_message=FakeMessage(text="bukan dokumen"))
        await bot.pdf_command(FakeUpdate(message), FakeContext({}))
        self.assertIn("Cara pakai /pdf", message.replies[0].text)

    # --- validasi tipe & ukuran

    async def test_unsupported_file_type(self):
        document = FakeDocument("foto.png", "image/png", 100)
        message = await self.run_reply_command("pdf", document, None)
        self.assertIn("bukan PDF", message.replies[0].text)
        self.assert_no_leftovers()

    async def test_word_command_with_pdf_is_rejected(self):
        document = FakeDocument("dokumen.pdf", PDF_MIME, 100)
        message = await self.run_reply_command("word", document, None)
        self.assertIn("bukan dokumen Word", message.replies[0].text)
        self.assert_no_leftovers()

    async def test_oversized_file_is_rejected(self):
        document = FakeDocument("besar.pdf", PDF_MIME, converter.CONVERT_MAX_FILE_SIZE + 1)
        message = await self.run_reply_command("pdf", document, None)
        self.assertIn("melebihi batas", message.replies[0].text)
        self.assert_no_leftovers()

    # --- alur sukses

    async def test_pdf_end_to_end(self):
        source = self.fixture("dokumen.pdf", make_pdf, text="Konten laporan tahunan")
        document = FakeDocument("dokumen.pdf", PDF_MIME, os.path.getsize(source), "pdf-file")
        message = await self.run_reply_command("pdf", document, source)

        sent = message.sent_documents
        self.assertEqual(len(sent), 1)
        self.assertTrue(sent[0].filename.endswith(".docx"))
        self.assertEqual(sent[0].data[:2], b"PK")  # docx adalah zip
        status = message.replies[0]
        self.assertTrue(status.deleted)  # pesan proses dihapus setelah selesai
        self.assert_no_leftovers()

    async def test_caption_command_is_supported(self):
        source = self.fixture("dokumen.pdf", make_pdf)
        document = FakeDocument("dokumen.pdf", PDF_MIME, os.path.getsize(source), "pdf-file")
        message = FakeMessage(caption="/pdf", document=document)
        await bot.convert_caption(FakeUpdate(message), FakeContext({document.file_id: source}))

        self.assertEqual(len(message.sent_documents), 1)
        self.assert_no_leftovers()

    async def test_caption_with_bot_username_is_supported(self):
        source = self.fixture("dokumen.pdf", make_pdf)
        document = FakeDocument("dokumen.pdf", PDF_MIME, os.path.getsize(source), "pdf-file")
        message = FakeMessage(caption="/pdf@my_bot", document=document)
        await bot.convert_caption(FakeUpdate(message), FakeContext({document.file_id: source}))
        self.assertEqual(len(message.sent_documents), 1)

    async def test_word_end_to_end_uses_ilovepdf(self):
        source = self.fixture("dokumen.docx", make_docx)
        document = FakeDocument("dokumen.docx", DOCX_MIME, os.path.getsize(source), "docx-file")

        def handler(request: httpx.Request) -> httpx.Response:
            path = request.url.path
            if path.startswith("/v1/start/"):
                return httpx.Response(200, json={"server": "mock.ilovepdf.test", "task": "task-1"})
            if path == "/v1/upload":
                return httpx.Response(200, json={"server_filename": "server-1.docx"})
            if path == "/v1/process":
                return httpx.Response(200, json={"status": "TaskSuccess"})
            if path == "/v1/download/task-1":
                return httpx.Response(
                    200, content=FAKE_PDF_BYTES, headers={"content-type": "application/pdf"}
                )
            if request.method == "DELETE" and path == "/v1/task/task-1":
                return httpx.Response(200, json={"status": "TaskDeleted"})
            return httpx.Response(404, json={"type": "HTTP Error"})

        transport = httpx.MockTransport(handler)

        with mock.patch.object(converter, "ILOVEPDF_PUBLIC_KEY", TEST_PUBLIC_KEY), mock.patch.object(
            converter, "ILOVEPDF_SECRET_KEY", TEST_SECRET_KEY
        ), patched_client(transport):
            message = await self.run_reply_command("word", document, source)

        sent = message.sent_documents
        self.assertEqual(len(sent), 1)
        self.assertTrue(sent[0].filename.endswith(".pdf"))
        self.assertEqual(sent[0].data, FAKE_PDF_BYTES)
        self.assertIn("iLovePDF", sent[0].caption_sent)
        self.assert_no_leftovers()

    async def test_status_message_edit_failure_does_not_fail_conversion(self):
        """Kegagalan edit/hapus pesan status tidak boleh menggagalkan konversi."""
        source = self.fixture("dokumen.pdf", make_pdf)
        document = FakeDocument("dokumen.pdf", PDF_MIME, os.path.getsize(source), "pdf-file")
        doc_message = FakeMessage(document=document)
        command_message = FakeMessage(text="/pdf", reply_to_message=doc_message)
        command_message.fail_status = True

        await bot.pdf_command(
            FakeUpdate(command_message), FakeContext({document.file_id: source})
        )

        self.assertEqual(len(command_message.sent_documents), 1)
        self.assert_no_leftovers()

    async def test_image_only_pdf_gets_ocr_warning(self):
        source = self.fixture("scan.pdf", make_image_only_pdf)
        document = FakeDocument("scan.pdf", PDF_MIME, os.path.getsize(source), "scan-file")
        message = await self.run_reply_command("pdf", document, source)

        sent = message.sent_documents
        self.assertEqual(len(sent), 1)
        self.assertIn("OCR", sent[0].caption_sent or "")

    # --- kegagalan

    async def test_corrupted_pdf_fails_gracefully(self):
        source = os.path.join(self.tmp, "rusak.pdf")
        with open(source, "wb") as fh:
            fh.write(os.urandom(128))
        document = FakeDocument("rusak.pdf", PDF_MIME, 128, "bad-file")
        message = await self.run_reply_command("pdf", document, source)

        self.assertEqual(message.sent_documents, [])
        status = message.replies[0]
        self.assertTrue(any("Gagal" in edit for edit in status.edits))
        self.assert_no_leftovers()

    async def test_password_protected_pdf_fails_gracefully(self):
        source = self.fixture("sandi.pdf", make_pdf, password="rahasia")
        document = FakeDocument("sandi.pdf", PDF_MIME, os.path.getsize(source), "lock-file")
        message = await self.run_reply_command("pdf", document, source)

        self.assertEqual(message.sent_documents, [])
        self.assertTrue(any("password" in edit for edit in message.replies[0].edits))
        self.assert_no_leftovers()

    async def test_empty_file_fails_gracefully(self):
        document = FakeDocument("kosong.pdf", PDF_MIME, 1, "empty-file")
        message = await self.run_reply_command("pdf", document, None)

        self.assertEqual(message.sent_documents, [])
        self.assertTrue(any("kosong" in edit for edit in message.replies[0].edits))
        self.assert_no_leftovers()

    async def test_word_without_api_keys_explains_setup(self):
        source = self.fixture("dokumen.docx", make_docx)
        document = FakeDocument("dokumen.docx", DOCX_MIME, os.path.getsize(source), "docx-file")

        with mock.patch.object(converter, "ILOVEPDF_PUBLIC_KEY", ""), mock.patch.object(
            converter, "ILOVEPDF_SECRET_KEY", ""
        ):
            message = await self.run_reply_command("word", document, source)

        self.assertEqual(message.sent_documents, [])
        status = message.replies[0]
        self.assertTrue(any("belum aktif" in edit for edit in status.edits))
        self.assert_no_leftovers()

    async def test_download_failure_is_reported_and_cleaned(self):
        document = FakeDocument("hilang.pdf", PDF_MIME, 10, "missing-file")
        message = await self.run_reply_command("pdf", document, None, files={})

        self.assertEqual(message.sent_documents, [])
        self.assertTrue(any("Gagal" in edit for edit in message.replies[0].edits))
        self.assert_no_leftovers()

    # --- duplikat & konkurensi

    async def test_duplicate_command_is_ignored_while_running(self):
        source = self.fixture("dokumen.pdf", make_pdf)
        document = FakeDocument("dokumen.pdf", PDF_MIME, os.path.getsize(source), "pdf-file")
        started = asyncio.Event()
        release = asyncio.Event()

        async def slow_convert(src_path, dst_path):
            started.set()
            await asyncio.wait_for(release.wait(), timeout=5)
            with open(dst_path, "wb") as fh:
                fh.write(b"PK\x03\x04 fake docx")

        async def second_run():
            doc_message = FakeMessage(document=document)
            command_message = FakeMessage(text="/pdf", reply_to_message=doc_message)
            await bot.pdf_command(
                FakeUpdate(command_message, 111), FakeContext({document.file_id: source})
            )
            return command_message

        with mock.patch.object(bot, "convert_pdf_to_word", slow_convert):
            first_message = FakeMessage(
                text="/pdf", reply_to_message=FakeMessage(document=document)
            )
            first_task = asyncio.create_task(
                bot.pdf_command(
                    FakeUpdate(first_message, 111), FakeContext({document.file_id: source})
                )
            )
            await asyncio.wait_for(started.wait(), timeout=5)

            second_message = await second_run()  # user yang sama, perintah dobel

            release.set()
            await asyncio.wait_for(first_task, timeout=10)

        self.assertEqual(len(first_message.sent_documents), 1)
        self.assertEqual(second_message.sent_documents, [])
        self.assertIn("masih diproses", second_message.replies[0].text.lower())
        self.assert_no_leftovers()

    async def test_multiple_users_convert_concurrently(self):
        sources, messages, contexts, updates = [], [], [], []
        for index in range(3):
            sources.append(self.fixture(f"user{index}.pdf", make_pdf, text=f"Dokumen pengguna {index}"))
        tasks = []
        for index, source in enumerate(sources):
            document = FakeDocument(f"user{index}.pdf", PDF_MIME, os.path.getsize(source), f"file-{index}")
            message = FakeMessage(text="/pdf", reply_to_message=FakeMessage(document=document))
            updates.append(message)
            contexts.append(FakeContext({document.file_id: source}))
            tasks.append(bot.pdf_command(FakeUpdate(message, 1000 + index), contexts[-1]))

        await asyncio.gather(*tasks)

        for message in updates:
            self.assertEqual(len(message.sent_documents), 1)
        self.assert_no_leftovers()


class ILovePDFClientTests(unittest.IsolatedAsyncioTestCase):
    """Alur API iLovePDF: autentikasi, sukses, error, cleanup."""

    def setUp(self):
        for target, value in (
            ("ILOVEPDF_PUBLIC_KEY", TEST_PUBLIC_KEY),
            ("ILOVEPDF_SECRET_KEY", TEST_SECRET_KEY),
            ("ILOVEPDF_REGION", ""),
        ):
            patcher = mock.patch.object(converter, target, value)
            patcher.start()
            self.addCleanup(patcher.stop)

    def make_source(self, name="dokumen.docx"):
        tmp = tempfile.mkdtemp(prefix="ilp_")
        self.addCleanup(shutil.rmtree, tmp, True)
        return make_docx(os.path.join(tmp, name))

    @staticmethod
    def use_transport(handler):
        return patched_client(httpx.MockTransport(handler))

    def test_token_is_valid_hs256_jwt(self):
        client = converter.ILovePDFClient()
        token = client._token()
        claims = pyjwt.decode(token, TEST_SECRET_KEY, algorithms=["HS256"])
        self.assertEqual(claims["jti"], TEST_PUBLIC_KEY)
        # iat = now - CLOCK_DELAY, exp = now + TOKEN_LIFETIME + CLOCK_DELAY (spt library resmi)
        self.assertEqual(
            claims["exp"] - claims["iat"],
            converter.TOKEN_LIFETIME + 2 * converter.CLOCK_DELAY,
        )
        header = pyjwt.get_unverified_header(token)
        self.assertEqual(header["alg"], "HS256")

    async def test_full_task_flow(self):
        calls = []

        def handler(request: httpx.Request) -> httpx.Response:
            calls.append((request.method, request.url.host, request.url.path))
            self.assertTrue(request.headers.get("authorization", "").startswith("Bearer "))
            path = request.url.path
            if path.startswith("/v1/start/"):
                return httpx.Response(
                    200, json={"server": "mock.ilovepdf.test", "task": "task-9", "remaining_credits": 10}
                )
            if path == "/v1/upload":
                self.assertIn(b'name="file"', request.content)
                return httpx.Response(200, json={"server_filename": "server-9.docx"})
            if path == "/v1/process":
                body = json.loads(request.content)
                self.assertEqual(body["tool"], "officepdf")
                self.assertEqual(body["task"], "task-9")
                self.assertEqual(body["files"][0]["server_filename"], "server-9.docx")
                return httpx.Response(200, json={"status": "TaskSuccess"})
            if path == "/v1/download/task-9":
                return httpx.Response(
                    200, content=FAKE_PDF_BYTES, headers={"content-type": "application/pdf"}
                )
            if request.method == "DELETE" and path == "/v1/task/task-9":
                return httpx.Response(200, json={"status": "TaskDeleted"})
            return httpx.Response(404, json={"type": "HTTP Error"})

        tmp = tempfile.mkdtemp(prefix="ilp_")
        self.addCleanup(shutil.rmtree, tmp, True)
        src = make_docx(os.path.join(tmp, "dokumen.docx"))
        dst = os.path.join(tmp, "dokumen.pdf")

        with self.use_transport(handler):
            await convert_word_to_pdf(src, dst, DOCX_MIME)

        with open(dst, "rb") as fh:
            self.assertEqual(fh.read(), FAKE_PDF_BYTES)

        paths = [path for _, _, path in calls]
        self.assertTrue(paths[0].startswith("/v1/start/officepdf"))
        self.assertIn("/v1/upload", paths)
        self.assertIn("/v1/process", paths)
        self.assertIn("/v1/download/task-9", paths)
        self.assertEqual(calls[-1][0], "DELETE")  # task dibersihkan terakhir
        self.assertIn("/v1/task/task-9", calls[-1][2])
        # upload, process, download, delete dilakukan di server worker
        self.assertTrue(all(host == "mock.ilovepdf.test" for _, host, _ in calls[1:]))

    async def test_task_is_deleted_even_when_process_fails(self):
        deleted = []

        def handler(request: httpx.Request) -> httpx.Response:
            path = request.url.path
            if path.startswith("/v1/start/"):
                return httpx.Response(200, json={"server": "mock.ilovepdf.test", "task": "task-x"})
            if path == "/v1/upload":
                return httpx.Response(200, json={"server_filename": "server-x.docx"})
            if path == "/v1/process":
                return httpx.Response(
                    400, json={"type": "ProcessingError", "message": "File is damaged"}
                )
            if request.method == "DELETE":
                deleted.append(path)
                return httpx.Response(200, json={"status": "TaskDeleted"})
            return httpx.Response(404)

        tmp = tempfile.mkdtemp(prefix="ilp_")
        self.addCleanup(shutil.rmtree, tmp, True)
        src = make_docx(os.path.join(tmp, "dokumen.docx"))

        with self.use_transport(handler):
            with self.assertRaises(ConversionError) as ctx:
                await convert_word_to_pdf(src, os.path.join(tmp, "out.pdf"), DOCX_MIME)

        self.assertIn("File is damaged", str(ctx.exception))
        self.assertEqual(deleted, ["/v1/task/task-x"])

    async def test_http_errors_are_translated(self):
        cases = [
            (401, "API key iLovePDF tidak valid"),
            (400, "menolak file ini"),
            (429, "Batas permintaan iLovePDF tercapai"),
            (500, "sedang bermasalah"),
        ]

        def make_handler(status):
            def handler(request: httpx.Request) -> httpx.Response:
                if request.url.path.startswith("/v1/start/"):
                    return httpx.Response(
                        status, json={"type": "HTTP Error", "message": "pesan uji"}
                    )
                return httpx.Response(404)

            return handler

        tmp = tempfile.mkdtemp(prefix="ilp_")
        self.addCleanup(shutil.rmtree, tmp, True)
        src = make_docx(os.path.join(tmp, "dokumen.docx"))

        for status, expected in cases:
            with self.subTest(status=status):
                with self.use_transport(make_handler(status)):
                    with self.assertRaises(ConversionError) as ctx:
                        await convert_word_to_pdf(src, os.path.join(tmp, f"out{status}.pdf"), DOCX_MIME)
                self.assertIn(expected, str(ctx.exception))

    async def test_timeout_is_translated(self):
        def handler(request: httpx.Request) -> httpx.Response:
            raise httpx.ReadTimeout("lambat")

        tmp = tempfile.mkdtemp(prefix="ilp_")
        self.addCleanup(shutil.rmtree, tmp, True)
        src = make_docx(os.path.join(tmp, "dokumen.docx"))

        with self.use_transport(handler):
            with self.assertRaises(ConversionError) as ctx:
                await convert_word_to_pdf(src, os.path.join(tmp, "out.pdf"), DOCX_MIME)

        self.assertIn("melebihi batas waktu", str(ctx.exception))

    async def test_missing_keys_raise_setup_hint(self):
        tmp = tempfile.mkdtemp(prefix="ilp_")
        self.addCleanup(shutil.rmtree, tmp, True)
        src = make_docx(os.path.join(tmp, "dokumen.docx"))

        with mock.patch.object(converter, "ILOVEPDF_PUBLIC_KEY", ""), mock.patch.object(
            converter, "ILOVEPDF_SECRET_KEY", ""
        ):
            with self.assertRaises(ConversionError) as ctx:
                await convert_word_to_pdf(src, os.path.join(tmp, "out.pdf"), DOCX_MIME)

        self.assertIn("belum aktif", str(ctx.exception))

    async def test_secret_key_is_never_logged(self):
        """Pastikan kunci API tidak pernah muncul di log."""
        records = []

        class Capture(logging.Handler):
            def emit(self, record):
                records.append(record.getMessage())

        capture = Capture()
        capture.setFormatter(logging.Formatter("%(message)s"))
        target_logger = logging.getLogger("converter")
        target_logger.addHandler(capture)
        self.addCleanup(target_logger.removeHandler, capture)

        def failing(request: httpx.Request) -> httpx.Response:
            return httpx.Response(401, json={"message": "unauthorized"})

        tmp = tempfile.mkdtemp(prefix="ilp_")
        self.addCleanup(shutil.rmtree, tmp, True)
        src = make_docx(os.path.join(tmp, "dokumen.docx"))

        with self.use_transport(failing):
            with self.assertRaises(ConversionError):
                await convert_word_to_pdf(src, os.path.join(tmp, "out.pdf"), DOCX_MIME)

        joined = "\n".join(records)
        self.assertTrue(records)  # ada log error, tapi tanpa rahasia
        self.assertNotIn(TEST_SECRET_KEY, joined)
        self.assertNotIn(TEST_PUBLIC_KEY, joined)


class RoutingTests(unittest.TestCase):
    """Pastikan update Telegram benar-benar diarahkan ke handler konversi."""

    TOKEN = "TEST_TOKEN_NOT_A_REAL_BOT"

    def setUp(self):
        from datetime import datetime, timezone

        from telegram import Bot, Chat, Message, Update, User

        self.bot_api = Bot(token=self.TOKEN)
        # username bot dibutuhkan CommandHandler; isi cache tanpa memanggil API
        self.bot_api._bot_user = User(
            id=1, is_bot=True, first_name="Test", username="test_bot"
        )
        self.chat = Chat(id=1, type="private")
        self.user = User(id=1, first_name="U", is_bot=False)
        self.Update = lambda message: Update(update_id=1, message=message)

        def make_message(**kwargs):
            message = Message(
                message_id=kwargs.pop("message_id", 1),
                date=datetime.now(timezone.utc),
                chat=self.chat,
                from_user=self.user,
                **kwargs,
            )
            message.set_bot(self.bot_api)  # dibutuhkan CommandHandler saat cek username bot
            return message

        self.Message = make_message

    def test_command_reply_matches_command_handler(self):
        from telegram import MessageEntity
        from telegram.ext import CommandHandler

        message = self.Message(
            message_id=2,
            text="/pdf",
            entities=[MessageEntity(type="bot_command", offset=0, length=4)],
        )
        handler = CommandHandler("pdf", bot.pdf_command)
        self.assertTrue(handler.check_update(self.Update(message)))

    def test_command_in_caption_matches_caption_filter(self):
        from telegram import Document
        from telegram.ext import filters

        document = Document(
            file_id="AAQ", file_unique_id="AQ", file_name="dok.pdf", mime_type="application/pdf"
        )
        message = self.Message(message_id=3, document=document, caption="/word")
        caption_filter = filters.Document.ALL & filters.CaptionRegex(
            r"^/(pdf|word)(?:@\w+)?(?:\s|$)"
        )
        self.assertTrue(caption_filter.check_update(self.Update(message)))

        # caption yang bukan perintah konversi tidak boleh tertangkap
        other = self.Message(message_id=4, document=document, caption="/download")
        self.assertFalse(caption_filter.check_update(self.Update(other)))

    def test_main_registers_conversion_handlers(self):
        from telegram.ext import Application

        captured = {}

        def fake_run(self, *args, **kwargs):
            captured["handlers"] = self.handlers

        with mock.patch.object(Application, "run_polling", fake_run):
            bot.main()

        callbacks = [
            getattr(handler, "callback", None).__name__
            for handler in captured["handlers"][0]
        ]
        for expected in ("pdf_command", "word_command", "convert_caption", "handle_message"):
            self.assertIn(expected, callbacks)
        # handler konversi harus terdaftar sebelum handler teks biasa
        self.assertLess(callbacks.index("convert_caption"), callbacks.index("handle_message"))


if __name__ == "__main__":
    unittest.main(verbosity=2)
